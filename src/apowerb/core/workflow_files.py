"""Contenu des fichiers produits par le nœud ``file`` d'un workflow.

Fonctions pures : une valeur déjà rendue entre, des octets sortent. Le
stockage (référence opaque du propriétaire) est fait ailleurs
(``workflow_runtime.services_for``) ; ``workflow_graph`` applique les bornes
et lève les codes produit.

Formats tabulaires (csv, xlsx) : une liste d'objets (colonnes = clés, dans
l'ordre de première apparition) ou une liste de listes dont la première est
l'en-tête ; une chaîne JSON de l'un ou l'autre est acceptée, puisqu'un agent
rend souvent du texte. Formats texte (markdown, text, docx, pdf) : une chaîne,
lue comme du Markdown simple pour docx et pdf (titres, listes, tableaux,
paragraphes) — pas un moteur de mise en page.
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Iterator, NamedTuple, Optional

# format -> (extension, type MIME)
FORMATS: dict[str, tuple[str, str]] = {
    "csv": (".csv", "text/csv"),
    "xlsx": (
        ".xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    "json": (".json", "application/json"),
    "markdown": (".md", "text/markdown"),
    "text": (".txt", "text/plain"),
    "docx": (
        ".docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    "pdf": (".pdf", "application/pdf"),
}
TABULAR = ("csv", "xlsx")
# Formats mis en page : leur rendu coûte au caractère (borné en amont).
RICH_TEXT = ("docx", "pdf")

# Une cellule qui commence ainsi est une formule pour Excel ou LibreOffice
# (injection CSV/XLSX) : elle est préfixée d'une apostrophe et reste du texte.
_FORMULA_START = ("=", "+", "@", "\t", "\r")
_NUMBER = re.compile(r"^-?\d+([.,]\d+)?$")
# Caractères de contrôle refusés par XML (xlsx, docx) : retirés.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class RowsError(ValueError):
    """La valeur ne se lit pas comme des lignes de tableau."""


def strip_control(text: str) -> str:
    """``text`` sans ses caractères de contrôle, tel que Word et PDF le rendent."""
    return _CONTROL.sub("", text)


def _as_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    return json.dumps(value, ensure_ascii=False, indent=2, default=str)


class TextTooLong(Exception):
    """Le texte à produire dépasse la borne demandée."""


def depth(value: Any, stop: int) -> int:
    """Profondeur d'imbrication de ``value`` (listes, objets), sans
    récursion ; le parcours s'arrête dès que ``stop`` est dépassé."""
    deepest = 0
    stack = [(value, 0)]
    while stack:
        item, level = stack.pop()
        if isinstance(item, dict):
            children = item.values()
        elif isinstance(item, list):
            children = item
        else:
            continue
        level += 1
        deepest = max(deepest, level)
        if deepest > stop:
            return deepest
        stack.extend((c, level) for c in children)
    return deepest


def bounded_text(value: Any, limit: int, as_json: bool = False) -> str:
    """Le texte de ``value`` (une valeur non textuelle en JSON indenté),
    produit morceau par morceau et abandonné dès que ``limit`` caractères
    sont dépassés (``TextTooLong``) : jamais la chaîne entière d'abord.
    ``as_json`` : tout passe par JSON (un texte devient une chaîne JSON)."""
    if isinstance(value, str) and not as_json:
        if len(value) > limit:
            raise TextTooLong
        return value
    if value is None and not as_json:
        return ""
    parts, size = [], 0
    encoder = json.JSONEncoder(ensure_ascii=False, indent=2, default=str)
    for chunk in encoder.iterencode(value):
        size += len(chunk)
        if size > limit:
            raise TextTooLong
        parts.append(chunk)
    return "".join(parts)


def json_value(value: Any) -> Any:
    """La valeur d'un fichier JSON : un texte JSON est relu, sinon gardé."""
    if isinstance(value, str):
        try:
            return json.loads(value)
        except (ValueError, RecursionError):
            return value
    return value


class Table(NamedTuple):
    """Valeur tabulaire lue sans être élargie : ses dimensions se bornent
    avant que la moindre cellule soit produite."""

    keys: list  # clés des objets, ou None pour des listes
    header: list[str]
    records: list  # lignes de données (objets ou listes)
    width: int  # nombre de colonnes

    def rows(self) -> Iterator[list[Any]]:
        if self.keys is None:
            return (list(r) for r in self.records)
        return ([row.get(k) for k in self.keys] for row in self.records)


def parse_table(value: Any) -> Table:
    """La ``Table`` d'une valeur tabulaire, ou ``RowsError``."""
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            raise RowsError("texte qui n'est pas une liste JSON") from None
    if isinstance(value, dict):
        # Un objet seul : une ligne. Un objet qui enveloppe la liste
        # (``{"rows": [...]}``) est trop ambigu pour être deviné.
        value = [value]
    if not isinstance(value, list) or not value:
        raise RowsError("liste de lignes vide ou absente")
    if all(isinstance(r, dict) for r in value):
        # dict : ordre de première apparition, en temps linéaire.
        keys = list(dict.fromkeys(k for row in value for k in row))
        return Table(keys, [str(k) for k in keys], value, len(keys))
    if all(isinstance(r, list) for r in value):
        width = max(len(r) for r in value)
        return Table(None, [str(c) for c in value[0]], value[1:], width)
    raise RowsError("lignes mêlant objets, listes ou valeurs simples")


def _cell(value: Any, neutralise: bool = True) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = _CONTROL.sub(
        "",
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, default=str),
    )
    if neutralise and (
        text.startswith(_FORMULA_START)
        or (text.startswith("-") and not _NUMBER.match(text))
    ):
        return "'" + text
    return text


def _csv(header: list[str], rows: Iterator[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow([_cell(h) for h in header])
    for row in rows:
        writer.writerow(["" if c is None else _cell(c) for c in row])
    # BOM : Excel ouvre alors les accents correctement.
    return buf.getvalue().encode("utf-8-sig")


def _xlsx(header: list[str], rows: Iterator[list[Any]], sheet: Optional[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    # Écriture en flux : la mémoire ne grandit pas avec le nombre de cellules.
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet or "Data")
    bold = Font(bold=True)

    def cells(values, header=False):
        out = []
        for value in values:
            # Pas d'apostrophe ici (elle resterait visible, ex. "+33…") : une
            # cellule de type texte n'est jamais évaluée comme formule.
            cell = WriteOnlyCell(ws, value=_cell(value, neutralise=False))
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if header:
                cell.font = bold
            out.append(cell)
        return out

    ws.append(cells(header, header=True))
    for row in rows:
        ws.append(cells(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_tables(text: str) -> tuple[int, int]:
    """(colonnes du plus large, cellules au total) des tableaux Markdown que
    ``_docx`` construirait : de quoi borner avant toute construction."""
    lines = text.splitlines()
    widest = cells = 0
    i = 0
    while i < len(lines):
        if (
            lines[i].lstrip().startswith("|")
            and i + 1 < len(lines)
            and _TABLE_SEP.match(lines[i + 1])
        ):
            width = len(_split_row(lines[i]))
            rows, i = 1, i + 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows, i = rows + 1, i + 1
            widest = max(widest, width)
            cells += rows * width
            continue
        i += 1
    return widest, cells


def _docx(text: str) -> bytes:
    import docx

    doc = docx.Document()
    lines = text.splitlines()
    i = 0
    paragraph: list[str] = []

    def flush():
        if paragraph:
            doc.add_paragraph(" ".join(paragraph))
            paragraph.clear()

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush()
        elif m := _HEADING.match(line):
            flush()
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif m := _BULLET.match(line):
            flush()
            doc.add_paragraph(m.group(1).strip(), style="List Bullet")
        elif m := _NUMBERED.match(line):
            flush()
            doc.add_paragraph(m.group(1).strip(), style="List Number")
        elif (
            line.lstrip().startswith("|")
            and i + 1 < len(lines)
            and _TABLE_SEP.match(lines[i + 1])
        ):
            flush()
            header = _split_row(line)
            body = []
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                body.append(_split_row(lines[i]))
                i += 1
            tbl = doc.add_table(rows=1 + len(body), cols=len(header))
            tbl.style = "Table Grid"
            for row_cells, values in zip(tbl.rows, [header, *body]):
                for cell, value in zip(row_cells.cells, values):
                    cell.text = value
            continue
        else:
            paragraph.append(line.strip())
        i += 1
    flush()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_font(pdf) -> str:
    """DejaVu Sans livrée avec matplotlib (dépendance du projet) : l'image
    de production n'installe aucune police système, et Helvetica refuse
    « € », « — » ou « œ »."""
    import os

    import matplotlib

    folder = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    styles = {
        "": "DejaVuSans.ttf",
        "B": "DejaVuSans-Bold.ttf",
        "I": "DejaVuSans-Oblique.ttf",
        "BI": "DejaVuSans-BoldOblique.ttf",
    }
    paths = {k: os.path.join(folder, v) for k, v in styles.items()}
    if all(os.path.exists(p) for p in paths.values()):
        for style, path in paths.items():
            pdf.add_font("DejaVuSans", style, path)
        return "DejaVuSans"
    from apowerb.core.agent_helpers.pdf_writer import _setup_pdf_font

    return _setup_pdf_font(pdf)


def _pdf(text: str) -> bytes:
    from fpdf import FPDF

    from apowerb.core.agent_helpers.text_utils import _md_to_html

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font(_pdf_font(pdf), size=11)
    pdf.write_html(_md_to_html(text))
    return bytes(pdf.output())


def build(
    fmt: str,
    value: Any,
    sheet: Optional[str] = None,
    table: Optional[Table] = None,
    text: Optional[str] = None,
) -> tuple[bytes, Optional[int]]:
    """(octets, nombre de lignes pour un format tabulaire) du fichier ``fmt``.
    ``table`` / ``text`` : la valeur déjà lue et bornée (``parse_table``,
    ``bounded_text``)."""
    if fmt in TABULAR:
        t = table or parse_table(value)
        rows = t.rows()
        data = _csv(t.header, rows) if fmt == "csv" else _xlsx(t.header, rows, sheet)
        return data, len(t.records)
    if fmt == "json":
        if text is None:
            text = json.dumps(
                json_value(value), ensure_ascii=False, indent=2, default=str
            )
        return text.encode(), None
    if text is None:
        text = _as_text(value)
    if fmt == "docx":
        return _docx(_CONTROL.sub("", text)), None
    if fmt == "pdf":
        return _pdf(_CONTROL.sub("", text)), None
    return text.encode("utf-8"), None
