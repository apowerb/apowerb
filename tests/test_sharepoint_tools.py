"""Tests for the SharePoint tools.

Graph request shapes are pinned against the v1.0 reference consulted
2026-10-06. httpx and the Microsoft token helper are mocked — no network, no
tenant. Word/PDF/Excel extraction runs on real in-memory documents.
"""

from __future__ import annotations

import io
from unittest.mock import MagicMock, patch

import httpx
import pytest

from apowerb.tools_store.portfolio import sharepoint
from apowerb.tools_store.portfolio.integration_status import (
    INTEGRATION_MISSING,
    IntegrationStatusError,
)

GRAPH = "https://graph.microsoft.com/v1.0"
SITE = "contoso.sharepoint.com,da60e844-ba1d-49bc-b4d4-d5e36bae9019,712a596e-90a1-49e3-9b48-bfa80bee8740"
DRIVE = "b!AbC-123_xyz"
ITEM = "01ABCDEF2345"


def _resp(payload: dict | None = None, content: bytes = b"") -> MagicMock:
    resp = MagicMock()
    resp.status_code = 200
    resp.json.return_value = payload or {}
    resp.content = content
    resp.raise_for_status.return_value = None
    return resp


@pytest.fixture(autouse=True)
def _token(monkeypatch):
    calls = []

    def fake_headers(prefix, *, scope, service_label):
        calls.append((prefix, scope))
        return {"Authorization": "Bearer graph-tok"}

    monkeypatch.setattr(sharepoint, "microsoft_auth_headers", fake_headers)
    return calls


def test_search_sites_uses_sharepoint_integration(_token):
    payload = {
        "value": [
            {"id": SITE, "displayName": "Ventes", "webUrl": "https://x/sites/ventes"}
        ]
    }
    with patch("httpx.get", return_value=_resp(payload)) as get:
        result = sharepoint.tool_sharepoint_search_sites("ventes")
    assert result["sites"][0] == {
        "id": SITE,
        "name": "Ventes",
        "webUrl": "https://x/sites/ventes",
        "description": None,
    }
    assert get.call_args.args[0] == f"{GRAPH}/sites"
    assert get.call_args.kwargs["params"] == {"search": "ventes"}
    assert get.call_args.kwargs["headers"]["Authorization"] == "Bearer graph-tok"
    assert _token == [("SHAREPOINT", "offline_access Sites.Read.All")]


def test_list_libraries_and_files_paths():
    with patch("httpx.get", return_value=_resp({"value": []})) as get:
        sharepoint.tool_sharepoint_list_libraries(SITE)
        sharepoint.tool_sharepoint_list_files(DRIVE)
        sharepoint.tool_sharepoint_list_files(DRIVE, folder_path="/Contrats/2026 Q1/")
        sharepoint.tool_sharepoint_list_files(DRIVE, folder_id=ITEM, limit=999)
    urls = [c.args[0] for c in get.call_args_list]
    assert urls == [
        f"{GRAPH}/sites/{SITE}/drives",
        f"{GRAPH}/drives/{DRIVE}/root/children",
        f"{GRAPH}/drives/{DRIVE}/root:/Contrats/2026%20Q1:/children",
        f"{GRAPH}/drives/{DRIVE}/items/{ITEM}/children",
    ]
    assert get.call_args.kwargs["params"] == {"$top": 200}


def test_search_files_escapes_quotes_and_keeps_drive_id():
    payload = {
        "value": [
            {
                "id": ITEM,
                "name": "Devis O'Neil.pdf",
                "file": {"mimeType": "application/pdf"},
                "parentReference": {"driveId": DRIVE},
            }
        ]
    }
    with patch("httpx.get", return_value=_resp(payload)) as get:
        result = sharepoint.tool_sharepoint_search_files("O'Neil devis", site_id=SITE)
    assert get.call_args.args[0] == (
        f"{GRAPH}/sites/{SITE}/drive/root/search(q='O%27%27Neil%20devis')"
    )
    assert result["items"][0]["driveId"] == DRIVE
    assert result["items"][0]["type"] == "file"


@pytest.mark.parametrize(
    "call",
    [
        lambda: sharepoint.tool_sharepoint_search_files("x"),
        lambda: sharepoint.tool_sharepoint_search_sites(" "),
        lambda: sharepoint.tool_sharepoint_list_libraries("../me/drive"),
        lambda: sharepoint.tool_sharepoint_read_file(DRIVE, "a/../b"),
    ],
)
def test_invalid_input_makes_no_request(call):
    with patch("httpx.get") as get:
        result = call()
    assert result["status"] == "error"
    get.assert_not_called()


def _docx_bytes(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=text)
    return bytes(pdf.output())


def _xlsx_bytes() -> bytes:
    import pandas as pd

    buf = io.BytesIO()
    pd.DataFrame({"client": ["Acme"], "ca": [1200]}).to_excel(buf, index=False)
    return buf.getvalue()


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("notes.txt", "Bonjour équipe".encode(), "Bonjour équipe"),
        (
            "contrat.docx",
            _docx_bytes("Clause de confidentialité"),
            "Clause de confidentialité",
        ),
        ("facture.pdf", _pdf_bytes("Total 1200 EUR"), "Total 1200 EUR"),
        ("ventes.xlsx", _xlsx_bytes(), "client,ca\nAcme,1200"),
    ],
)
def test_read_file_extracts_text(name, content, expected):
    meta = {"id": ITEM, "name": name, "size": len(content), "file": {"mimeType": "x"}}
    with patch("httpx.get", side_effect=[_resp(meta), _resp(content=content)]) as get:
        result = sharepoint.tool_sharepoint_read_file(DRIVE, ITEM)
    assert result["status"] == "success"
    assert expected in result["content"]
    assert result["truncated"] is False
    assert get.call_args.args[0] == f"{GRAPH}/drives/{DRIVE}/items/{ITEM}/content"
    assert get.call_args.kwargs["follow_redirects"] is True


def test_read_file_unsupported_and_too_large():
    meta = {
        "id": ITEM,
        "name": "logo.png",
        "size": 10,
        "file": {},
        "webUrl": "https://x/logo",
    }
    with patch("httpx.get", side_effect=[_resp(meta), _resp(content=b"\x89PNG")]):
        result = sharepoint.tool_sharepoint_read_file(DRIVE, ITEM)
    assert result["content"] is None
    assert result["webUrl"] == "https://x/logo"

    big = {"id": ITEM, "name": "big.pdf", "size": 50 * 1024 * 1024, "file": {}}
    with patch("httpx.get", return_value=_resp(big)) as get:
        result = sharepoint.tool_sharepoint_read_file(DRIVE, ITEM)
    assert result["content"] is None
    assert get.call_count == 1  # no download


def test_missing_integration_is_reported(monkeypatch):
    def missing(*_a, **_k):
        raise IntegrationStatusError(
            INTEGRATION_MISSING, "microsoft_sharepoint", "Not connected"
        )

    monkeypatch.setattr(sharepoint, "microsoft_auth_headers", missing)
    result = sharepoint.tool_sharepoint_search_sites("ventes")
    assert result["status"] == "integration_status"
    assert result["code"] == INTEGRATION_MISSING


def test_tenant_block_423_mapped():
    request = httpx.Request("GET", f"{GRAPH}/sites")
    err = httpx.HTTPStatusError(
        "locked", request=request, response=httpx.Response(423, request=request)
    )
    with patch("httpx.get", side_effect=err):
        result = sharepoint.tool_sharepoint_search_sites("ventes")
    assert result["http_status"] == 423
    assert "admin center" in result["message"]


def test_discovered_as_oauth_category():
    from apowerb.tools_store.tool_manager import ToolsStore, _category_requires_oauth

    tools = ToolsStore().get_tools_in_category("sharepoint")
    assert "sharepoint.tool_sharepoint_read_file" in tools
    assert len(tools) == 5
    assert _category_requires_oauth("sharepoint") is True
