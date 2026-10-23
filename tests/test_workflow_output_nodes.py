"""Nœuds de sortie ``file`` et ``file_save``, pièces jointes du mail (moteur seul).

Le stockage du fichier, l'envoi vers OneDrive et la résolution des pièces
jointes sont simulés via ``NodeServices`` et ``run_notify`` ; leur
branchement réel (propriétaire contrôlé, bornes) est couvert par
``test_workflow_output_services.py``.
"""

import asyncio
import io
import json

import pytest

from apowerb.core import workflow_graph as wg

T = {"id": "t", "type": "trigger"}
ROWS = [
    {"mois": "2026-09", "ventes": 120, "region": "Nord"},
    {"mois": "2026-10", "ventes": 135.5, "region": "Sud"},
]


class FakeServices:
    def __init__(self):
        self.saved = []
        self.drive = []

    async def save_file(self, name, content_type, data):
        ref = f"att_{len(self.saved):032x}"
        self.saved.append(
            {"name": name, "content_type": content_type, "data": data, "ref": ref}
        )
        return {
            "name": name,
            "ref": ref,
            "size": len(data),
            "content_type": content_type,
        }

    async def save_to_drive(self, provider, ref, folder, conflict):
        self.drive.append((provider, ref, folder, conflict))
        return {
            "name": "ventes.csv",
            "path": f"{folder}/ventes.csv",
            "web_url": "https://drive.example/ventes.csv",
        }

    def bundle(self):
        return wg.NodeServices(
            save_file=self.save_file, save_to_drive=self.save_to_drive
        )


def _graph(nodes, edges):
    return wg.WorkflowGraph.model_validate(
        {"version": 1, "nodes": nodes, "edges": edges}
    )


def _chain(*nodes):
    nodes = [T, *nodes]
    edges = [{"source": a["id"], "target": b["id"]} for a, b in zip(nodes, nodes[1:])]
    return nodes, edges


def _run(nodes, edges, payload=None, services=None, run_notify=None):
    async def run_tool(tool, args):
        pytest.fail("pas d'outil")

    async def run_agent(agent_id, message):
        pytest.fail("pas d'agent")

    async def go():
        out = []
        async for chunk in wg.run_graph(
            _graph(nodes, edges),
            payload=payload or {},
            run_agent=run_agent,
            run_tool=run_tool,
            cancel_event=asyncio.Event(),
            services=services.bundle() if services else None,
            run_notify=run_notify,
        ):
            out.append(json.loads(chunk[len("data: ") :]))
        return out

    return asyncio.run(go())


def _complete(events, node_id):
    return next(
        e for e in events if e["event"] == "node_complete" and e["node_id"] == node_id
    )


def _error(events, node_id):
    return next(
        e for e in events if e["event"] == "node_error" and e["node_id"] == node_id
    )


def _file(cfg, id_="f"):
    return {"id": id_, "type": "file", "config": cfg}


def _invalid(node):
    with pytest.raises(wg.GraphError) as exc:
        wg.validate_node(wg.Node.model_validate(node))
    return exc.value


# --- validation du nœud file ---------------------------------------------------


def test_file_format_must_be_known():
    err = _invalid(_file({"format": "exe", "content": "x"}))
    assert err.code == "file_format_unknown"


def test_file_content_is_required():
    err = _invalid(_file({"format": "csv", "content": "  "}))
    assert err.code == "file_content_missing"


def test_file_sheet_name_follows_excel_rules():
    err = _invalid(
        _file({"format": "xlsx", "content": "{{t.rows}}", "sheet_name": "a/b"})
    )
    assert err.code == "file_sheet_invalid"
    err = _invalid(
        _file({"format": "xlsx", "content": "{{t.rows}}", "sheet_name": "x" * 32})
    )
    assert err.code == "file_sheet_invalid"


def test_file_filename_is_bounded():
    err = _invalid(
        _file({"format": "csv", "content": "{{t.rows}}", "filename": "x" * 121})
    )
    assert err.code == "file_filename_invalid"


def test_every_documented_format_validates():
    for fmt in wg.FILE_FORMATS:
        wg.validate_node(
            wg.Node.model_validate(_file({"format": fmt, "content": "{{t.rows}}"}))
        )


# --- exécution du nœud file ----------------------------------------------------


def test_csv_from_rows_is_stored_and_described():
    svc = FakeServices()
    nodes, edges = _chain(
        _file(
            {"format": "csv", "content": "{{t.rows}}", "filename": "ventes-{{t.mois}}"}
        )
    )
    events = _run(nodes, edges, {"rows": ROWS, "mois": "oct"}, svc)
    out = _complete(events, "f")["output"]
    assert set(out) == {"name", "ref", "size", "content_type", "format", "rows"}
    assert (
        out["name"] == "ventes-oct.csv" and out["format"] == "csv" and out["rows"] == 2
    )
    saved = svc.saved[0]
    assert saved["content_type"] == "text/csv"
    text = saved["data"].decode("utf-8-sig")
    assert text.splitlines()[0] == "mois,ventes,region"
    assert "2026-10,135.5,Sud" in text
    assert out["size"] == len(saved["data"])


def test_file_appears_in_node_and_run_artifacts():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": ROWS}, svc)
    art = _complete(events, "f")["artifacts"][0]
    assert art["file_ref"] == svc.saved[0]["ref"] and art["file_name"] == "f.csv"
    done = next(e for e in events if e["event"] == "done")
    assert done["artifacts"][0]["file_ref"] == svc.saved[0]["ref"]
    assert done["artifacts"][0]["node_id"] == "f"


def test_csv_accepts_rows_given_as_a_json_string():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.text}}"}))
    _run(nodes, edges, {"text": json.dumps(ROWS)}, svc)
    assert svc.saved[0]["data"].decode("utf-8-sig").startswith("mois,ventes,region")


def test_csv_accepts_a_list_of_lists_with_a_header_row():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    _run(nodes, edges, {"rows": [["a", "b"], [1, 2]]}, svc)
    assert svc.saved[0]["data"].decode("utf-8-sig").splitlines() == ["a,b", "1,2"]


def test_csv_refuses_plain_text():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.text}}"}))
    events = _run(nodes, edges, {"text": "bonjour"}, svc)
    assert _error(events, "f")["code"] == "file_bad_rows"
    assert svc.saved == []


def test_too_many_rows_is_refused(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_ROWS", 1)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": ROWS}, svc)
    assert _error(events, "f")["code"] == "file_too_many_rows"


def test_file_over_the_size_cap_is_never_stored(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_BYTES", 10)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "text", "content": "{{t.text}}"}))
    events = _run(nodes, edges, {"text": "x" * 50}, svc)
    assert _error(events, "f")["code"] == "file_too_large"
    assert svc.saved == []


def test_xlsx_has_the_header_and_the_rows():
    import openpyxl

    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "xlsx", "content": "{{t.rows}}", "sheet_name": "Ventes"})
    )
    _run(nodes, edges, {"rows": ROWS}, svc)
    saved = svc.saved[0]
    assert saved["name"] == "f.xlsx"
    assert (
        saved["content_type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    ws = openpyxl.load_workbook(io.BytesIO(saved["data"]))["Ventes"]
    assert [c.value for c in ws[1]] == ["mois", "ventes", "region"]
    assert [c.value for c in ws[3]] == ["2026-10", 135.5, "Sud"]


def test_xlsx_neutralises_formulas():
    import openpyxl

    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "xlsx", "content": "{{t.rows}}"}))
    rows = [{"a": '=HYPERLINK("http://x")'}, {"a": "+33612345678"}]
    _run(nodes, edges, {"rows": rows}, svc)
    ws = openpyxl.load_workbook(io.BytesIO(svc.saved[0]["data"])).active
    # Une cellule texte n'est jamais évaluée : la valeur reste intacte.
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == '=HYPERLINK("http://x")'
    assert ws["A3"].value == "+33612345678"


def test_csv_neutralises_formulas():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    rows = [{"=cmd": "=1+1"}, {"=cmd": "-5"}, {"=cmd": "-1+1"}, {"=cmd": "-2,5"}]
    _run(nodes, edges, {"rows": rows}, svc)
    lines = svc.saved[0]["data"].decode("utf-8-sig").splitlines()
    # L'en-tête aussi ; un nombre négatif reste un nombre.
    assert lines == ["'=cmd", "'=1+1", "-5", "'-1+1", '"-2,5"']


# --- bornes avant construction ---------------------------------------------------


def _no_build(monkeypatch):
    def boom(*a, **k):
        pytest.fail("rien ne doit être construit au-delà des bornes")

    monkeypatch.setattr(wg.workflow_files, "build", boom)


def test_heterogeneous_keys_are_bounded_before_building(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_COLUMNS", 5)
    _no_build(monkeypatch)
    svc = FakeServices()
    rows = [{f"k{r}_{i}": 1 for i in range(2)} for r in range(3)]  # 6 colonnes
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": rows}, svc)
    err = _error(events, "f")
    assert err["code"] == "file_too_many_columns"
    assert err["params"] == {"node": "f", "max": "5"}
    assert svc.saved == []


def test_lists_count_their_widest_row_as_columns(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_COLUMNS", 2)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": "xlsx", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": [["a"], [1, 2, 3]]}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_columns"


def test_cells_are_bounded_before_building(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_CELLS", 5)
    _no_build(monkeypatch)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "xlsx", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": ROWS}, svc)  # 2 lignes x 3 colonnes
    err = _error(events, "f")
    assert err["code"] == "file_too_many_cells"
    assert err["params"] == {"node": "f", "max": "5"}
    assert svc.saved == []


def test_table_is_read_once_outside_the_event_loop(monkeypatch):
    import threading

    seen = []
    real = wg.workflow_files.parse_table

    def spy(value):
        seen.append(threading.current_thread() is threading.main_thread())
        return real(value)

    monkeypatch.setattr(wg.workflow_files, "parse_table", spy)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    _run(nodes, edges, {"rows": ROWS}, svc)
    assert seen == [False]
    assert svc.saved


def test_heterogeneous_header_keeps_first_seen_order():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    _run(nodes, edges, {"rows": [{"b": 1, "a": 2}, {"c": 3, "a": 4}]}, svc)
    lines = svc.saved[0]["data"].decode("utf-8-sig").splitlines()
    assert lines == ["b,a,c", "1,2,", ",4,3"]


def test_render_failure_logs_the_error_type_only(monkeypatch, caplog):
    def fail(*a, **k):
        raise ValueError("cellule secrète 4242")

    monkeypatch.setattr(wg.workflow_files, "build", fail)
    nodes, edges = _chain(_file({"format": "text", "content": "x"}))
    with caplog.at_level("WARNING"):
        events = _run(nodes, edges, {}, FakeServices())
    assert _error(events, "f")["code"] == "file_build_failed"
    assert "ValueError" in caplog.text
    assert "4242" not in caplog.text


def test_control_characters_do_not_break_xlsx_or_docx():
    import docx
    import openpyxl

    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "xlsx", "content": "{{t.rows}}"}),
        _file({"format": "docx", "content": "{{t.text}}"}, id_="g"),
    )
    events = _run(nodes, edges, {"rows": [{"a": "x\x01y"}], "text": "a\x0bb"}, svc)
    assert _complete(events, "g")
    ws = openpyxl.load_workbook(io.BytesIO(svc.saved[0]["data"])).active
    assert ws["A2"].value == "xy"
    doc = docx.Document(io.BytesIO(svc.saved[1]["data"]))
    assert [p.text for p in doc.paragraphs if p.text] == ["ab"]


def test_empty_content_is_refused_clearly():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "text", "content": "{{t.text}}"}))
    events = _run(nodes, edges, {"text": ""}, svc)
    assert _error(events, "f")["code"] == "file_empty"
    assert svc.saved == []


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_long_text_is_refused_before_rendering(monkeypatch, fmt):
    monkeypatch.setattr(wg, "MAX_FILE_TEXT_CHARS", 10)

    def boom(*a, **k):
        pytest.fail("aucun rendu au-delà de la borne")

    monkeypatch.setattr(wg.workflow_files, "build", boom)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": fmt, "content": "{{t.text}}"}))
    events = _run(nodes, edges, {"text": "x" * 50}, svc)
    err = _error(events, "f")
    assert err["code"] == "file_text_too_long"
    assert err["params"] == {"node": "f", "format": fmt, "max": "10"}
    assert svc.saved == []


def test_plain_text_formats_are_only_bound_by_size(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_TEXT_CHARS", 10)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "text", "content": "{{t.text}}"}))
    _run(nodes, edges, {"text": "x" * 50}, svc)
    assert svc.saved[0]["data"] == b"x" * 50


def test_json_is_indented():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "json", "content": "{{t.rows}}"}))
    _run(nodes, edges, {"rows": ROWS}, svc)
    assert json.loads(svc.saved[0]["data"]) == ROWS
    assert svc.saved[0]["content_type"] == "application/json"


def test_markdown_keeps_the_text():
    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "markdown", "content": "# Rapport\n{{t.text}}"})
    )
    _run(nodes, edges, {"text": "Hausse de 12 %"}, svc)
    assert svc.saved[0]["data"].decode() == "# Rapport\nHausse de 12 %"
    assert svc.saved[0]["name"] == "f.md"


def test_docx_renders_headings_bullets_and_tables():
    import docx

    svc = FakeServices()
    report = "# Prévisions\n\nTexte é.\n\n- point un\n- point deux\n\n| mois | ventes |\n|---|---|\n| oct | 135 |"
    nodes, edges = _chain(_file({"format": "docx", "content": report}))
    _run(nodes, edges, {}, svc)
    doc = docx.Document(io.BytesIO(svc.saved[0]["data"]))
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Heading 1", "Prévisions") in styles
    assert ("List Bullet", "point deux") in styles
    table = doc.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["oct", "135"]


def test_pdf_is_a_pdf():
    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "pdf", "content": "# Rapport\nVentes : 135 € — été"})
    )
    _run(nodes, edges, {}, svc)
    assert svc.saved[0]["data"].startswith(b"%PDF")
    assert svc.saved[0]["content_type"] == "application/pdf"


def test_filename_cannot_escape_and_keeps_its_extension():
    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "csv", "content": "{{t.rows}}", "filename": "{{t.name}}"})
    )
    _run(nodes, edges, {"rows": ROWS, "name": "../../etc/passwd"}, svc)
    name = svc.saved[0]["name"]
    assert "/" not in name and ".." not in name and name.endswith(".csv")


def test_file_without_storage_service_fails_clearly():
    nodes, edges = _chain(_file({"format": "csv", "content": "{{t.rows}}"}))
    events = _run(nodes, edges, {"rows": ROWS}, None)
    err = _error(events, "f")
    assert err["code"] == "service_unavailable"


def test_file_fields_are_checked_statically():
    nodes, edges = _chain(
        _file({"format": "csv", "content": "{{t.rows}}"}),
        {"id": "o", "type": "output", "config": {"value": "{{f.nope}}"}},
    )
    with pytest.raises(wg.GraphError):
        wg.validate_graph(_graph(nodes, edges))
    nodes[-1]["config"]["value"] = "{{f.ref}} {{f.name}} {{f.size}}"
    wg.validate_graph(_graph(nodes, edges))


# --- nœud file_save ------------------------------------------------------------


def _save(cfg, id_="s"):
    return {"id": id_, "type": "file_save", "config": cfg}


def test_file_save_provider_must_be_known():
    err = _invalid(_save({"provider": "dropbox", "file": "{{f}}"}))
    assert err.code == "file_save_provider_unknown"


def test_file_save_needs_a_file():
    err = _invalid(_save({"provider": "onedrive", "file": ""}))
    assert err.code == "file_save_file_missing"


def test_file_save_folder_cannot_climb():
    err = _invalid(_save({"provider": "onedrive", "file": "{{f}}", "folder": "a/../b"}))
    assert err.code == "file_save_folder_invalid"


def test_file_save_conflict_is_checked():
    err = _invalid(
        _save({"provider": "onedrive", "file": "{{f}}", "conflict": "merge"})
    )
    assert err.code == "file_save_conflict_unknown"


def test_file_save_empty_conflict_means_rename():
    node = _save({"provider": "onedrive", "file": "{{t.ref}}", "conflict": None})
    wg.validate_node(wg.Node.model_validate(node))
    svc = FakeServices()
    nodes, edges = _chain(node)
    _run(nodes, edges, {"ref": "att_" + "a" * 32}, svc)
    assert svc.drive == [("onedrive", "att_" + "a" * 32, "", "rename")]


def test_file_save_sends_the_produced_file():
    svc = FakeServices()
    nodes, edges = _chain(
        _file({"format": "csv", "content": "{{t.rows}}"}),
        _save(
            {"provider": "onedrive", "file": "{{f}}", "folder": "Rapports/{{t.mois}}"}
        ),
    )
    events = _run(nodes, edges, {"rows": ROWS, "mois": "oct"}, svc)
    assert svc.drive == [("onedrive", svc.saved[0]["ref"], "Rapports/oct", "rename")]
    out = _complete(events, "s")["output"]
    assert out == {
        "provider": "onedrive",
        "name": "ventes.csv",
        "path": "Rapports/oct/ventes.csv",
        "web_url": "https://drive.example/ventes.csv",
    }


def test_file_save_accepts_a_bare_ref():
    svc = FakeServices()
    nodes, edges = _chain(
        _save({"provider": "onedrive", "file": "{{t.ref}}", "conflict": "replace"})
    )
    _run(nodes, edges, {"ref": "att_" + "a" * 32}, svc)
    assert svc.drive == [("onedrive", "att_" + "a" * 32, "", "replace")]


def test_file_save_refuses_something_that_is_not_a_file():
    svc = FakeServices()
    nodes, edges = _chain(_save({"provider": "onedrive", "file": "{{t.x}}"}))
    events = _run(nodes, edges, {"x": {"nom": "pas un fichier"}}, svc)
    assert _error(events, "s")["code"] == "file_save_bad_file"
    assert svc.drive == []


def test_rendered_folder_cannot_climb_either():
    svc = FakeServices()
    nodes, edges = _chain(
        _save({"provider": "onedrive", "file": "{{t.ref}}", "folder": "{{t.dir}}"})
    )
    events = _run(nodes, edges, {"ref": "att_" + "a" * 32, "dir": "../secret"}, svc)
    assert _error(events, "s")["code"] == "file_save_folder_invalid"
    assert svc.drive == []


# --- pièces jointes du mail ----------------------------------------------------


def _mail(cfg):
    base = {
        "channel": "email",
        "to": ["a@example.test"],
        "subject": "Rapport",
        "body": "Ci-joint",
    }
    return {"id": "n", "type": "notification", "config": {**base, **cfg}}


def test_attachments_are_for_email_only():
    node = {
        "id": "n",
        "type": "notification",
        "config": {"channel": "teams", "subject": "x", "attachments": "{{f}}"},
    }
    err = _invalid(node)
    assert err.code == "notification_attachments_email_only"


def test_email_receives_the_produced_files():
    svc = FakeServices()
    calls = []

    async def run_notify(node_id, channel, to, subject, body, attachments=None):
        calls.append(attachments)
        return len(to)

    nodes, edges = _chain(
        _file({"format": "csv", "content": "{{t.rows}}"}),
        _mail({"attachments": "{{f}}"}),
    )
    events = _run(nodes, edges, {"rows": ROWS}, svc, run_notify)
    assert calls == [[svc.saved[0]["ref"]]]
    # Même forme de sortie qu'un mail sans pièce jointe.
    assert _complete(events, "n")["output"] == {"channel": "email", "sent": 1}


def test_email_without_attachments_keeps_the_old_call():
    calls = []

    async def run_notify(node_id, channel, to, subject, body):
        calls.append(subject)
        return 1

    nodes, edges = _chain(_mail({}))
    events = _run(nodes, edges, {}, None, run_notify)
    assert calls == ["Rapport"]
    assert _complete(events, "n")["output"] == {"channel": "email", "sent": 1}


def test_attachments_that_render_to_nothing_are_refused():
    async def run_notify(*a, **k):
        pytest.fail("rien ne doit partir")

    nodes, edges = _chain(_mail({"attachments": "{{t.refs}}"}))
    events = _run(nodes, edges, {"refs": []}, None, run_notify)
    assert _error(events, "n")["code"] == "notification_attachment_missing"


def test_too_many_attachments_are_refused(monkeypatch):
    monkeypatch.setattr(wg, "MAX_NOTIFICATION_ATTACHMENTS", 1)

    async def run_notify(*a, **k):
        pytest.fail("rien ne doit partir")

    nodes, edges = _chain(_mail({"attachments": "{{t.refs}}"}))
    events = _run(
        nodes, edges, {"refs": ["att_" + "a" * 32, "att_" + "b" * 32]}, None, run_notify
    )
    assert _error(events, "n")["code"] == "notification_too_many_attachments"


# --- texte : bornes avant sérialisation et avant rendu ---------------------------


def _deep(depth, leaf=1):
    value = leaf
    for _ in range(depth):
        value = [value]
    return value


@pytest.mark.parametrize("fmt", ["json", "text", "markdown", "pdf"])
def test_deep_values_are_refused_before_serialising(monkeypatch, fmt):
    monkeypatch.setattr(wg, "MAX_FILE_DEPTH", 5)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": fmt, "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": _deep(6)}, FakeServices())
    err = _error(events, "f")
    assert err["code"] == "file_too_deep"
    assert err["params"] == {"node": "f", "max": "5"}


def test_json_given_as_text_is_held_to_the_same_depth(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_DEPTH", 5)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": "json", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": json.dumps(_deep(6))}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_deep"


def test_json_stops_serialising_at_the_size_cap(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_BYTES", 100)
    _no_build(monkeypatch)
    chunks = []
    real = json.JSONEncoder.iterencode

    def spy(self, o, _one_shot=False):
        for chunk in real(self, o, _one_shot):
            chunks.append(chunk)
            yield chunk

    monkeypatch.setattr(json.JSONEncoder, "iterencode", spy)
    nodes, edges = _chain(_file({"format": "json", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": list(range(5000))}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_large"
    # Arrêt à la borne : la sérialisation complète ferait ~5 000 morceaux.
    assert len(chunks) < 200


def test_rich_text_from_a_structure_is_measured_without_building_it(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_TEXT_CHARS", 50)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": "pdf", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": {"k": list(range(5000))}}, FakeServices())
    assert _error(events, "f")["code"] == "file_text_too_long"


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_markdown_tables_are_bounded_in_cells(monkeypatch, fmt):
    monkeypatch.setattr(wg, "MAX_DOC_TABLE_CELLS", 5)
    _no_build(monkeypatch)
    text = "| a | b | c |\n|---|---|---|\n| 1 | 2 | 3 |"
    nodes, edges = _chain(_file({"format": fmt, "content": text}))
    events = _run(nodes, edges, {}, FakeServices())
    err = _error(events, "f")
    assert err["code"] == "file_too_many_cells"
    assert err["params"] == {"node": "f", "max": "5"}


def test_markdown_tables_are_bounded_in_columns(monkeypatch):
    monkeypatch.setattr(wg, "MAX_DOC_TABLE_COLUMNS", 2)
    _no_build(monkeypatch)
    text = "| a | b | c |\n|---|---|---|"
    nodes, edges = _chain(_file({"format": "docx", "content": text}))
    events = _run(nodes, edges, {}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_columns"


def test_cells_of_all_tables_add_up(monkeypatch):
    monkeypatch.setattr(wg, "MAX_DOC_TABLE_CELLS", 7)
    _no_build(monkeypatch)
    table = "| a | b |\n|---|---|\n| 1 | 2 |"  # 4 cellules
    nodes, edges = _chain(_file({"format": "docx", "content": table + "\n\n" + table}))
    events = _run(nodes, edges, {}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_cells"


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_rich_text_is_bounded_in_lines(monkeypatch, fmt):
    monkeypatch.setattr(wg, "MAX_FILE_TEXT_LINES", 3)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": fmt, "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": "a\nb\nc\nd"}, FakeServices())
    err = _error(events, "f")
    assert err["code"] == "file_too_many_lines"
    assert err["params"] == {"node": "f", "max": "3"}


def test_a_small_word_table_is_still_rendered():
    import docx

    svc = FakeServices()
    text = "| a | b |\n|---|---|\n| 1 | 2 |"
    nodes, edges = _chain(_file({"format": "docx", "content": text}))
    _run(nodes, edges, {}, svc)
    doc = docx.Document(io.BytesIO(svc.saved[0]["data"]))
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["1", "2"]


def test_errors_never_quote_the_content():
    nodes, edges = _chain(_save({"provider": "onedrive", "file": "{{t.x}}"}))
    events = _run(nodes, edges, {"x": {"nom": "secret-4242"}}, FakeServices())
    assert _error(events, "s")["code"] == "file_save_bad_file"
    assert "4242" not in json.dumps([e for e in events if "error" in e["event"]])

    async def run_notify(*a, **k):
        pytest.fail("rien ne doit partir")

    nodes, edges = _chain(_mail({"attachments": "{{t.x}}"}))
    events = _run(nodes, edges, {"x": [{"nom": "secret-4242"}]}, None, run_notify)
    assert _error(events, "n")["code"] == "notification_attachment_bad_ref"
    assert "4242" not in json.dumps([e for e in events if "error" in e["event"]])


def test_plain_text_in_a_json_file_stays_valid_json():
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "json", "content": "{{t.v}}"}))
    _run(nodes, edges, {"v": "bonjour"}, svc)
    assert json.loads(svc.saved[0]["data"]) == "bonjour"


# --- caractères de contrôle et fins de ligne : bornes comptées sur le texte rendu ---


@pytest.mark.parametrize("ctrl", ["\x1c", "\x0b", "\x0c", "\x1f"])
def test_a_control_character_in_the_separator_does_not_hide_a_table(monkeypatch, ctrl):
    monkeypatch.setattr(wg, "MAX_DOC_TABLE_CELLS", 5)
    _no_build(monkeypatch)
    text = f"| a | b | c |\n|-{ctrl}-|---|---|\n| 1 | 2 | 3 |"
    nodes, edges = _chain(_file({"format": "docx", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": text}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_cells"


def test_carriage_returns_count_as_lines(monkeypatch):
    monkeypatch.setattr(wg, "MAX_FILE_TEXT_LINES", 3)
    _no_build(monkeypatch)
    nodes, edges = _chain(_file({"format": "docx", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": "- a\r- b\r- c\r- d"}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_lines"


def test_the_widest_table_counts_even_when_it_comes_second(monkeypatch):
    monkeypatch.setattr(wg, "MAX_DOC_TABLE_COLUMNS", 2)
    _no_build(monkeypatch)
    narrow = "| a | b |\n|---|---|"
    wide = "| a | b | c |\n|---|---|---|"
    nodes, edges = _chain(_file({"format": "docx", "content": "{{t.v}}"}))
    events = _run(nodes, edges, {"v": narrow + "\n\n" + wide}, FakeServices())
    assert _error(events, "f")["code"] == "file_too_many_columns"


def test_the_text_checked_is_the_text_rendered(monkeypatch):
    seen = {}
    real = wg.workflow_files.build

    def spy(fmt, value, sheet=None, table=None, text=None):
        seen["text"] = text
        return real(fmt, value, sheet, table=table, text=text)

    monkeypatch.setattr(wg.workflow_files, "build", spy)
    svc = FakeServices()
    nodes, edges = _chain(_file({"format": "docx", "content": "{{t.v}}"}))
    _run(nodes, edges, {"v": "titre\x1c suite"}, svc)
    assert seen["text"] == "titre suite"
